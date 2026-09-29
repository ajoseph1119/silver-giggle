from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class SpreadsheetService:

    HEADERS = [
        "Title",
        "Company",
        "Location",
        "URL",
        "Snippet",
        "Source",
    ]

    def create(self, jobs: list[dict], filename: str = "jobs.xlsx"):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Jobs"

        # Headers
        for column, header in enumerate(self.HEADERS, start=1):
            cell = worksheet.cell(
                row=1,
                column=column,
                value=header
            )

            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

        # Job results
        for row, job in enumerate(jobs, start=2):
            worksheet.cell(row=row, column=1, value=job["title"])
            worksheet.cell(row=row, column=2, value=job["company"])
            worksheet.cell(row=row, column=3, value=job["location"])
            worksheet.cell(row=row, column=4, value=job["url"])
            worksheet.cell(row=row, column=5, value=job["snippet"])
            worksheet.cell(row=row, column=6, value=job["source"])

        # Make columns readable
        widths = {
            1: 35,
            2: 25,
            3: 25,
            4: 60,
            5: 80,
            6: 25,
        }

        for column, width in widths.items():
            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = width

        # Wrap long text
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        # Freeze header
        worksheet.freeze_panes = "A2"

        # Enable filtering
        worksheet.auto_filter.ref = worksheet.dimensions

        workbook.save(filename)

        print(f"Spreadsheet created: {filename}")